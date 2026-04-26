import requests
import datetime
from collections import Counter
import re
import openpyxl

def get_news():
    url = 'https://belarusbank.by/api/news_info'
    params = {'lang': 'ru'}
    response = requests.get(url, params=params)
    data = response.json()
    return data[:20]


def is_odd_day(date_str):
    """Проверяет, является ли день месяца нечётным."""
    try:
        dt = datetime.datetime.strptime(date_str, "%Y-%m-%d %H:%M:%S")
    except ValueError:
        try:
            dt = datetime.datetime.strptime(date_str, "%Y-%m-%d")
        except ValueError:
            return False
    return dt.day % 2 != 0


def filter_news_by_odd_day(news_list):
    return [news for news in news_list if is_odd_day(news.get("start_date", ""))]


def analyze_news(news_list):
    # Самый длинный текст
    longest_text_news = max(news_list, key=lambda x: len(x.get("html_ru", "")))
    # Самый длинный заголовок (по словам)
    longest_title_news = max(news_list, key=lambda x: len(x.get("name_ru", "").split()))
    # Самая частая русская буква
    all_text = " ".join(news.get("html_ru", "") for news in news_list)
    russian_letters = re.findall(r'[а-яё]', all_text.lower())
    if russian_letters:
        letter_counts = Counter(russian_letters)
        most_common_letter = letter_counts.most_common(1)[0]
    else:
        most_common_letter = ("нет русских букв", 0)
    return {
        "longest_text": {
            "title": longest_text_news.get("name_ru", ""),
            "length": len(longest_text_news.get("html_ru", "")),
            "date": longest_text_news.get("start_date", ""),
            "text": longest_text_news.get("html_ru", "")[:200] + "..."
        },
        "most_common_russian_letter": most_common_letter,
        "longest_title": {
            "title": longest_title_news.get("name_ru", ""),
            "word_count": len(longest_title_news.get("name_ru", "").split()),
            "date": longest_title_news.get("start_date", "")
        }
    }


def save_to_excel(all_news, filtered_news, analysis_results, filename="news_analysis.xlsx"):
    wb = openpyxl.Workbook()

    # Лист с исходными 20 новостями
    ws_orig = wb.active
    ws_orig.title = "Исходные 20 новостей"
    ws_orig.append(["Заголовок (ru)", "Дата", "Текст (первые 200 символов)"])
    for news in all_news:
        title = news.get("name_ru", "")
        date = news.get("start_date", "")
        snippet = news.get("html_ru", "")[:200]
        if len(news.get("html_ru", "")) > 200:
            snippet += "..."
        ws_orig.append([title, date, snippet])

    # Лист с отфильтрованными новостями
    ws_filt = wb.create_sheet("Отфильтрованные (нечётные дни)")
    ws_filt.append(["Заголовок (ru)", "Дата", "Текст (полный)"])
    for news in filtered_news:
        ws_filt.append([news.get("name_ru", ""), news.get("start_date", ""), news.get("html_ru", "")])

    # Лист с результатами анализа
    ws_analysis = wb.create_sheet("Результаты анализа")
    ws_analysis.append(["Параметр", "Значение", "Дополнительно"])
    if analysis_results:
        lt = analysis_results["longest_text"]
        ws_analysis.append(
            ["Самый длинный текст", f"{lt['length']} символов", f"Заголовок: {lt['title']}, Дата: {lt['date']}"])
        ws_analysis.append(["Фрагмент текста:", lt['text'], ""])

        letter, count = analysis_results["most_common_russian_letter"]
        ws_analysis.append(["Самая частая русская буква", f"'{letter}'", f"Встречается {count} раз(а)"])

        ltitle = analysis_results["longest_title"]
        ws_analysis.append(["Самый длинный заголовок", f"{ltitle['word_count']} слов",
                            f"Заголовок: {ltitle['title']}, Дата: {ltitle['date']}"])

    # Автоподбор ширины колонок
    for ws in [ws_orig, ws_filt, ws_analysis]:
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    max_len = max(max_len, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    wb.save(filename)
    print(f"Результаты сохранены в {filename}")


def main():
    print("1. Получение 20 последних новостей...")
    all_news = get_news()
    if not all_news:
        print("Не удалось получить новости.")
        return
    print(f"Получено новостей: {len(all_news)}")

    print("2. Фильтрация по нечётным дням...")
    filtered_news = filter_news_by_odd_day(all_news)
    print(f"Осталось после фильтрации: {len(filtered_news)}")

    print("3. Анализ текстов...")
    analysis_results = analyze_news(filtered_news)

    print("4. Сохранение в Excel...")
    save_to_excel(all_news, filtered_news, analysis_results)
    print("Готово!")


if __name__ == "__main__":
    main()